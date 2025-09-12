from openpyxl.styles import Alignment

class workbook():

    def __init__(self, wb, categories, block_lists) -> None:

        
        self.block_lists         = block_lists
        self.content_categories  = categories
        self.workbook            = wb
    
    def singleSiteOneNetwork(self, data):
        
        ws = self.workbook['Network']

        ws["B8"]   = data['customerName']     
        ws["B12"]  = data["address"]
        ws["A17"]  = data['size']
        ws["B17"]  = data["hostname"]
        ws["C17"]  = "No"
        ws["D17"]  = "hub"
        ws["E17"]  = "yes"
        ws["F17"]  = "static"
        ws["A25"]  = data["hostname"]
        ws["B25"]  = "WAN1"
        ws["C25"]  = data["provider"]
        ws["D25"]  = data["ipAssignment"]
        ws["E25"]  = data["network"]
        ws["F25"]  = data["gateway"]
        ws['G25']  = data["description"]
        ws["A49"]  = data["hostname"]
        ws['B49']  = "LAN1"
        ws["C49"]  = "data"
        ws["D49"]  = "1"
        ws["E49"]  = "10.0.0.0/24"
        ws["F49"]  = "10.0.0.1"
        ws["G49"]  = "internal"
        ws["A58"]  = data["hostname"]
        ws["B58"]  = "1"
        ws["C58"]  = "8.8.8.8"
        ws["D58"]  = "75.75.75.75"
        ws["G58"]  = "10.0.0.2-10.0.0.240"
        ws["A66"]  = data["hostname"]
        ws["B66"]  = "1"
        ws["C66"]  = "10.0.0.1"

    def moderateContentFilterSBS(self, business_type):

        ws = self.workbook["Content Categories"]
        
        block_list = self.block_lists.get(business_type, "general business")
        
            
        for k, v in self.content_categories.items():

            ws[f'F{v}'].alignment = Alignment(horizontal="center", vertical="center")
            
            ws[f'F{v}'] = "Y" if k in block_list else "N"

